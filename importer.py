"""
Excel (.xlsx) ва Word (.docx) файлҳоро хонда, ба SQLite илова мекунад.

Формати Excel / Word (динамикӣ — аз DB месозад):
  Сутун A : Адабӣ           (калимаи адабӣ — ҳатмӣ)
  Сутун B…: Ном-и-ноҳия     (ҳар ноҳияе ки дар DB мавҷуд аст)
  Охирин-1: Навъ             (калима / ҷумла — ихтиёрӣ)
  Охирин  : Ҷузъи нутқ      (исм/сифат/феъл/… — ихтиёрӣ)
"""
import re
import db


# ══════════════════════════════════════════════════════════════════════
# Ёрдамчии мутобиқсозии ном → калид
# ══════════════════════════════════════════════════════════════════════

def _name_to_key(name: str, dialect_map: dict[str, str]) -> str | None:
    """
    Номи сарлавҳаро ба dialect_key мувофиқат медиҳад.
    dialect_map: {key: name}
    """
    nl = name.strip().lower()
    for key, dname in dialect_map.items():
        if dname.lower() == nl:
            return key
    # Ҷустуҷӯи қисман (мисол: "Ғармӣ" дар "Лаҳҷаи Ғармӣ")
    for key, dname in dialect_map.items():
        if nl in dname.lower() or dname.lower() in nl:
            return key
    return None


# ══════════════════════════════════════════════════════════════════════
# Асос: татбиқи сатрҳо ба DB
# ══════════════════════════════════════════════════════════════════════

def _apply_rows(rows) -> tuple[int, int]:
    """rows = list of (literary, {dk: form}, kind, pos)"""
    added = 0
    skipped = 0
    for literary, translations, kind, pos in rows:
        if not literary or not str(literary).strip():
            skipped += 1
            continue
        cat = "phrases" if str(kind).strip().lower() in (
            "ҷумла", "jumla", "phrase") else "words"
        for dk, form in translations.items():
            if form and str(form).strip():
                status = db.add_word(dk, str(literary).strip(),
                                     str(form).strip(), cat, pos or "")
                if status in ("added", "updated"):
                    added += 1
    return added, skipped


# ══════════════════════════════════════════════════════════════════════
# Excel
# ══════════════════════════════════════════════════════════════════════

def import_excel(path: str) -> tuple[int, int, str]:
    try:
        import openpyxl
    except ImportError:
        return 0, 0, "openpyxl насб нашудааст (pip install openpyxl)"

    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        all_rows = list(ws.iter_rows(values_only=True))
        wb.close()
    except Exception as e:
        return 0, 0, f"Файл кушода нашуд: {e}"

    if not all_rows:
        return 0, 0, "Файл холӣ аст"

    # Сарлавҳа аз сатри якум
    header = [str(c).strip() if c else "" for c in all_rows[0]]
    dialect_map = {d["key"]: d["name"] for d in db.get_dialects()}

    # col_index → dialect_key  (сутун 0 = адабӣ, рад мешавад)
    col_dk: dict[int, str] = {}
    kind_col = pos_col = -1
    for i, h in enumerate(header):
        if i == 0:
            continue
        hl = h.lower()
        if hl in ("навъ", "nav", "type"):
            kind_col = i
        elif hl in ("ҷузъи нутқ", "pos", "part of speech"):
            pos_col = i
        else:
            dk = _name_to_key(h, dialect_map)
            if dk:
                col_dk[i] = dk

    if not col_dk:
        return 0, 0, ("Сутунҳои лаҳҷавӣ ёфт нашуданд.\n"
                      "Сарлавҳаи сутунҳо бояд бо номи ноҳияҳои луғат мувофиқ бошанд.")

    rows = []
    for row in all_rows[1:]:
        if not row or not row[0]:
            continue
        literary     = str(row[0]).strip()
        translations = {}
        for ci, dk in col_dk.items():
            val = row[ci] if ci < len(row) else None
            if val:
                translations[dk] = str(val).strip()
        kind = str(row[kind_col]).strip() if kind_col >= 0 and kind_col < len(row) and row[kind_col] else "калима"
        pos  = str(row[pos_col]).strip()  if pos_col  >= 0 and pos_col  < len(row) and row[pos_col]  else ""
        rows.append((literary, translations, kind, pos))

    added, skipped = _apply_rows(rows)
    return added, skipped, ""


# ══════════════════════════════════════════════════════════════════════
# Word
# ══════════════════════════════════════════════════════════════════════

def import_word(path: str) -> tuple[int, int, str]:
    try:
        from docx import Document
    except ImportError:
        return 0, 0, "python-docx насб нашудааст (pip install python-docx)"

    try:
        doc = Document(path)
    except Exception as e:
        return 0, 0, f"Файл кушода нашуд: {e}"

    if not doc.tables:
        return 0, 0, "Дар файли Word ҷадвал ёфт нашуд"

    dialect_map = {d["key"]: d["name"] for d in db.get_dialects()}
    rows = []

    for table in doc.tables:
        if not table.rows:
            continue
        # Сарлавҳа
        header = [c.text.strip() for c in table.rows[0].cells]
        col_dk: dict[int, str] = {}
        kind_col = pos_col = -1
        for i, h in enumerate(header):
            if i == 0:
                continue
            hl = h.lower()
            if hl in ("навъ", "nav", "type"):
                kind_col = i
            elif hl in ("ҷузъи нутқ", "pos", "part of speech"):
                pos_col = i
            else:
                dk = _name_to_key(h, dialect_map)
                if dk:
                    col_dk[i] = dk

        for row in table.rows[1:]:
            cells = [c.text.strip() for c in row.cells]
            if not cells or not cells[0]:
                continue
            literary     = cells[0]
            translations = {}
            for ci, dk in col_dk.items():
                val = cells[ci] if ci < len(cells) else ""
                if val:
                    translations[dk] = val
            kind = cells[kind_col] if kind_col >= 0 and kind_col < len(cells) else "калима"
            pos  = cells[pos_col]  if pos_col  >= 0 and pos_col  < len(cells) else ""
            rows.append((literary, translations, kind, pos))

    if not rows:
        return 0, 0, "Сутунҳои лаҳҷавӣ ёфт нашуданд"

    added, skipped = _apply_rows(rows)
    return added, skipped, ""


# ══════════════════════════════════════════════════════════════════════
# Шаблони Excel — динамикӣ аз DB
# ══════════════════════════════════════════════════════════════════════

def create_template_excel(save_path: str) -> str:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        return "openpyxl насб нашудааст"

    dialects = db.get_dialects()
    if not dialects:
        return "Луғат холӣ аст — аввал ноҳия илова кунед"

    wb  = openpyxl.Workbook()
    ws  = wb.active
    ws.title = "Калимаҳо"

    # ── Сарлавҳаҳо ──────────────────────────────────────────────────
    headers = ["Адабӣ"] + [d["name"] for d in dialects] + ["Навъ", "Ҷузъи нутқ"]

    # Ранги сарлавҳа: адабӣ=кабуди торик, ноҳияҳо=рангҳои гуногун, охирон=хокистарӣ
    import colorsys
    n = len(dialects)
    dial_colors = []
    for i in range(n):
        h = i / max(n, 1)
        r, g, b = colorsys.hsv_to_rgb(h, 0.7, 0.55)
        dial_colors.append(f"{int(r*255):02x}{int(g*255):02x}{int(b*255):02x}")

    hdr_colors = ["1a3a6e"] + dial_colors + ["3a3a3a", "2a4a3a"]

    thin   = Side(style="thin", color="555555")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, (h, hc) in enumerate(zip(headers, hdr_colors), start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font      = Font(bold=True, color="FFFFFF", name="Segoe UI", size=10)
        cell.fill      = PatternFill("solid", fgColor=hc)
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
        cell.border    = border

    ws.row_dimensions[1].height = 28

    # ── Намунаҳо (3 сатр) ────────────────────────────────────────────
    examples_lit = ["ман меравам", "намедонам", "хона"]
    for r_i, lit in enumerate(examples_lit, start=2):
        cell = ws.cell(row=r_i, column=1, value=lit)
        cell.font      = Font(name="Segoe UI", size=10, color="DDDDDD", italic=True)
        cell.fill      = PatternFill("solid", fgColor="1a1d27")
        cell.alignment = Alignment(vertical="center")
        cell.border    = border
        for c_i in range(2, len(headers) + 1):
            cell2 = ws.cell(row=r_i, column=c_i, value="")
            cell2.font   = Font(name="Segoe UI", size=10, color="AAAAAA")
            cell2.fill   = PatternFill("solid", fgColor="1a1d27")
            cell2.border = border

    # Сатри холии ҷудокунак
    sep_row = len(examples_lit) + 2
    for c_i in range(1, len(headers) + 1):
        cell = ws.cell(row=sep_row, column=c_i, value="")
        cell.fill   = PatternFill("solid", fgColor="2a2a2a")
        cell.border = border

    # Сатри дастур
    hint_row = sep_row + 1
    cell = ws.cell(row=hint_row, column=1,
                   value="← Аз ин сатр калимаҳои худро нависед")
    cell.font      = Font(name="Segoe UI", size=9, color="FFD54F", italic=True)
    cell.fill      = PatternFill("solid", fgColor="1a1a2a")
    cell.border    = border
    for c_i in range(2, len(headers) + 1):
        cell2 = ws.cell(row=hint_row, column=c_i, value="")
        cell2.fill   = PatternFill("solid", fgColor="1a1a2a")
        cell2.border = border

    # Сатрҳои холии воридот (20 сатр)
    for r_i in range(hint_row + 1, hint_row + 21):
        for c_i in range(1, len(headers) + 1):
            cell = ws.cell(row=r_i, column=c_i, value="")
            cell.font   = Font(name="Segoe UI", size=10)
            cell.border = border
            cell.alignment = Alignment(vertical="center")

    # ── Паҳно ────────────────────────────────────────────────────────
    ws.column_dimensions["A"].width = 26
    for i in range(2, len(dialects) + 2):
        ws.column_dimensions[get_column_letter(i)].width = 18
    ws.column_dimensions[get_column_letter(len(headers) - 1)].width = 12
    ws.column_dimensions[get_column_letter(len(headers))].width     = 14

    # Яхдон кардани сатри якум
    ws.freeze_panes = "A2"

    # ── Варақи дастур ────────────────────────────────────────────────
    ws2 = wb.create_sheet("Дастур")
    lines = [
        "ДАСТУР — БАРОИ ПУРКУНИИ ҶАДВАЛ",
        "",
        "Сутун А  →  Калима ё ҷумлаи АДАБИИ тоҷикӣ (ҳатмӣ)",
    ]
    for i, d in enumerate(dialects):
        col_letter = get_column_letter(i + 2)
        lines.append(f"Сутун {col_letter}  →  Тарҷума ба «{d['name']}» ({d.get('region','—')})")
    n_cols = len(dialects) + 2
    lines += [
        f"Сутун {get_column_letter(n_cols)}    →  Навъ: «калима» ё «ҷумла» (ихтиёрӣ)",
        f"Сутун {get_column_letter(n_cols+1)}  →  Ҷузъи нутқ: исм/сифат/феъл/зарф (ихтиёрӣ)",
        "",
        "Агар лаҳҷаро надонед — хонаро холӣ гузоред.",
        "Сатри якум (сарлавҳа) дигар накунед.",
        "Файлро ҳамчун .xlsx захира кунед.",
        "",
        "Ноҳияҳои дар луғат мавҷуд:",
    ]
    for d in dialects:
        lines.append(f"  • {d['name']}  ({d.get('region','—')})")

    for i, line in enumerate(lines, start=1):
        cell = ws2.cell(row=i, column=1, value=line)
        is_hdr = (i == 1)
        is_dial = line.startswith("  •")
        cell.font = Font(
            name="Segoe UI", size=11 if is_hdr else 10,
            bold=is_hdr,
            color="FFD54F" if is_hdr else ("90EE90" if is_dial else "CCCCCC"))
        cell.fill = PatternFill("solid", fgColor="0f1117")
    ws2.column_dimensions["A"].width = 62

    try:
        wb.save(save_path)
        return ""
    except Exception as e:
        return str(e)
